import json
from datetime import datetime, timezone
from io import BytesIO
from types import SimpleNamespace

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.reports.cli import main as report_cli_main
from app.reports.repository import get_report_repository, reset_report_repository
from app.reports.service import _confirmation_secret, reset_report_service
from app.reports.schemas import ReportArtifact
from app.reports.storage import ArtifactStore
from app.reports.workflow import reset_workflow_service
from scripts import rollback_report_storage_bucket

AUTH_HEADER = {"Authorization": "Bearer report-test-token"}
CALLBACK_HEADER = {"X-Callback-Token": "callback-test-token"}


#: Marcado pela #228 (S4-32B): este arquivo cruza fronteira entre
#: componentes (API, modelo, banco ou servico externo). Rodar so estes com
#: `pytest -m integration`.
pytestmark = pytest.mark.integration


@pytest.fixture(autouse=True)
def report_environment(monkeypatch, tmp_path):
    identities = {
        "report-test-token": {
            "subject": "user-report-1",
            "display_name": "Usuario de Teste",
            "role": "pmo",
            "can_export": True,
            "allow_all": True,
        },
        "restricted-token": {
            "subject": "user-restricted",
            "display_name": "Usuario Restrito",
            "role": "equipe_projeto",
            "can_export": True,
            "allow_all": False,
            "projects": ["prj-0442"],
            "subportfolios": [],
        },
    }
    monkeypatch.setenv("REPORT_SIMULATED_IDENTITIES", json.dumps(identities))
    monkeypatch.setenv("REPORT_CALLBACK_TOKEN", "callback-test-token")
    monkeypatch.setenv("REPORT_CONFIRMATION_SECRET", "confirmation-test-secret")
    monkeypatch.setenv("REPORT_REPOSITORY_MODE", "memory")
    monkeypatch.setenv("REPORT_EXECUTION_MODE", "local")
    monkeypatch.setenv("REPORT_STORAGE_MODE", "local")
    monkeypatch.setenv("REPORT_PORTFOLIO_SOURCE", "synthetic")
    monkeypatch.setenv("PLN_WARMUP", "0")
    monkeypatch.setenv("REPORT_ARTIFACTS_DIR", str(tmp_path / "reports"))
    _reset_runtime()
    yield
    _reset_runtime()


def _reset_runtime():
    reset_report_service()
    reset_report_repository()
    reset_workflow_service()
    _confirmation_secret.cache_clear()


def _prepare_payload(scope_id="sp-modernizacao", scope_type="subportfolio"):
    return {
        "contract_version": "1.0",
        "conversation_id": "conversation-001",
        "interaction_id": "interaction-001",
        "report_type": "executive_status",
        "scope": {"type": scope_type, "id": scope_id},
        "period": {
            "start_date": "2026-08-01",
            "end_date": "2026-08-31",
            "label": "Agosto de 2026",
        },
        "options": {
            "audience": "executive_board",
            "purpose": "monthly_governance",
            "emphasis": ["schedule", "risks", "benefits"],
            "additional_context": "Destacar decisoes do proximo ciclo.",
            "output_format": "docx",
        },
    }


def _prepare(client, headers=None, payload=None):
    return client.post(
        "/api/report-requests/prepare",
        json=payload or _prepare_payload(),
        headers=headers or AUTH_HEADER,
    )


def _create_payload(prepared):
    return {
        "prepared_request": prepared["prepared_request"],
        "confirmation_hash": prepared["confirmation_hash"],
        "idempotency_key": "idem-report-001",
    }


def test_prepare_requires_authorized_scope():
    with TestClient(app) as client:
        response = _prepare(
            client,
            headers={"Authorization": "Bearer restricted-token"},
        )

    assert response.status_code == 403
    assert response.json()["error"]["code"] == "REPORT_ACCESS_DENIED"


def test_prepare_returns_confirmable_versioned_contract():
    with TestClient(app) as client:
        response = _prepare(client)

    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "ready_for_confirmation"
    assert len(data["confirmation_hash"]) == 64
    assert data["prepared_request"]["requested_by"] == "user-report-1"
    assert data["prepared_request"]["scope"]["id"] == "sp-modernizacao"
    assert "Modernização da Rede" in data["confirmation_summary"]


def test_local_execution_generates_downloadable_docx_and_is_idempotent():
    with TestClient(app) as client:
        prepared = _prepare(client).json()
        payload = _create_payload(prepared)
        created = client.post("/api/report-requests", json=payload, headers=AUTH_HEADER)

        assert created.status_code == 202
        execution_id = created.json()["execution_id"]

        current = client.get(
            f"/api/report-requests/{execution_id}", headers=AUTH_HEADER
        )
        assert current.status_code == 200
        assert current.json()["status"] == "completed_with_warnings"
        assert current.json()["artifact"]["sha256"]
        assert "storage_reference" not in current.text

        downloaded = client.get(
            f"/api/report-requests/{execution_id}/artifact", headers=AUTH_HEADER
        )
        assert downloaded.status_code == 200
        assert downloaded.content.startswith(b"PK")
        document = Document(BytesIO(downloaded.content))
        document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        assert "Relatório executivo" in document_text
        assert "Modernização da Rede" in document_text
        assert "Fontes, atualização e limitações" in document_text
        assert execution_id in document_text

        repeated = client.post(
            "/api/report-requests", json=payload, headers=AUTH_HEADER
        )
        assert repeated.status_code == 200
        assert repeated.json()["execution_id"] == execution_id


def test_confirmation_hash_rejects_mutated_parameters():
    with TestClient(app) as client:
        prepared = _prepare(client).json()
        payload = _create_payload(prepared)
        payload["prepared_request"]["scope"]["id"] = "prj-0442"
        response = client.post(
            "/api/report-requests", json=payload, headers=AUTH_HEADER
        )

    assert response.status_code == 409
    assert response.json()["error"]["code"] == "REPORT_CONFLICT"


def test_power_automate_callback_is_authenticated_and_idempotent(monkeypatch):
    monkeypatch.setenv("REPORT_EXECUTION_MODE", "power_automate")
    monkeypatch.setattr(
        "app.reports.service.ReportService.dispatch_power_automate",
        lambda _service, _execution_id: None,
    )
    _reset_runtime()
    with TestClient(app) as client:
        prepared = _prepare(client).json()
        created = client.post(
            "/api/report-requests",
            json=_create_payload(prepared),
            headers=AUTH_HEADER,
        )
        execution = created.json()
        callback = {
            "contract_version": "1.0",
            "callback_id": "callback-report-001",
            "execution_id": execution["execution_id"],
            "correlation_id": execution["correlation_id"],
            "flow_run_id": "flow-run-001",
            "status": "completed",
            "result": {"schema_version": "1.0"},
        }

        denied = client.post(
            "/api/integrations/power-automate/report-results",
            json=callback,
            headers={"X-Callback-Token": "wrong"},
        )
        accepted = client.post(
            "/api/integrations/power-automate/report-results",
            json=callback,
            headers=CALLBACK_HEADER,
        )
        repeated = client.post(
            "/api/integrations/power-automate/report-results",
            json=callback,
            headers=CALLBACK_HEADER,
        )
        conflicting_callback = {**callback, "flow_run_id": "flow-run-different"}
        conflicted = client.post(
            "/api/integrations/power-automate/report-results",
            json=conflicting_callback,
            headers=CALLBACK_HEADER,
        )

    assert denied.status_code == 401
    assert denied.json()["error"]["code"] == "REPORT_CALLBACK_UNAUTHORIZED"
    assert accepted.status_code == 200
    assert accepted.json()["status"] == "completed"
    assert get_report_repository().callback_payload_hash(callback["callback_id"])
    assert repeated.status_code == 200
    assert repeated.json()["execution_id"] == execution["execution_id"]
    assert conflicted.status_code == 409
    assert conflicted.json()["error"]["code"] == "REPORT_CONFLICT"


def test_power_automate_callback_rejects_invalid_payload():
    with TestClient(app) as client:
        response = client.post(
            "/api/integrations/power-automate/report-results",
            headers=CALLBACK_HEADER,
            json={"callback_id": "short", "unexpected": True},
        )

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "REPORT_INVALID_INPUT"


def test_power_automate_callback_rejects_unknown_execution():
    with TestClient(app) as client:
        response = client.post(
            "/api/integrations/power-automate/report-results",
            headers=CALLBACK_HEADER,
            json={
                "contract_version": "1.0",
                "callback_id": "callback-report-unknown",
                "execution_id": "bc436aed-b4fc-443c-ab5f-7564c89664a1",
                "correlation_id": "correlation-unknown",
                "flow_run_id": "flow-run-unknown",
                "status": "completed",
                "result": {"schema_version": "1.0"},
            },
        )

    assert response.status_code == 404
    assert response.json()["error"]["code"] == "REPORT_NOT_FOUND"


def test_power_automate_simulator_exercises_callback_contract(monkeypatch):
    monkeypatch.setenv("REPORT_EXECUTION_MODE", "power_automate_simulator")
    _reset_runtime()
    with TestClient(app) as client:
        prepared = _prepare(client).json()
        created = client.post(
            "/api/report-requests",
            json=_create_payload(prepared),
            headers=AUTH_HEADER,
        )
        execution_id = created.json()["execution_id"]
        current = client.get(
            f"/api/report-requests/{execution_id}", headers=AUTH_HEADER
        )

    assert created.status_code == 202
    assert current.status_code == 200
    assert current.json()["status"] == "completed_with_warnings"
    assert current.json()["artifact"]["storage_provider"] == "local"
    assert any("simulador local" in item for item in current.json()["warnings"])


def test_report_intent_is_dispatched_without_generative_provider(monkeypatch):
    classification = SimpleNamespace(
        macrointencao="TRANSACAO",
        intencao="transacao_gerar_relatorio",
        confianca_macro=0.99,
        confianca_intencao=0.98,
        fluxo="Workflow guiado de relatorio",
        handler="handler:transacao_gerar_relatorio",
        alternativas=[],
        roteada=True,
        resposta_sugerida=None,
    )
    monkeypatch.setattr("app.nlp.inference.classify", lambda _text: classification)
    monkeypatch.setattr(
        "app.pipeline.generate_chat_response",
        lambda *_args, **_kwargs: pytest.fail(
            "Provider generativo nao deveria ser chamado"
        ),
    )

    with TestClient(app) as client:
        response = client.post(
            "/api/chat",
            json={
                "message": "gere um relatorio executivo",
                "use_intent_classifier": True,
            },
        )

    assert response.status_code == 200
    data = response.json()
    assert data["model"] == "workflow-router"
    assert data["workflow"]["workflow"] == "transacao_gerar_relatorio"
    assert data["intent"]["handler"] == "handler:transacao_gerar_relatorio"


def test_invalid_report_input_uses_controlled_error_envelope():
    payload = _prepare_payload()
    payload["period"]["end_date"] = "2025-01-01"
    with TestClient(app) as client:
        response = _prepare(client, payload=payload)

    assert response.status_code == 422
    assert response.json() == {
        "error": {
            "code": "REPORT_INVALID_INPUT",
            "message": "Revise os parametros informados para o relatorio.",
        }
    }


def test_guided_workflow_runs_from_questions_to_downloadable_artifact():
    with TestClient(app) as client:
        started = client.post(
            "/api/workflow-sessions",
            headers=AUTH_HEADER,
            json={
                "contract_version": "1.0",
                "conversation_id": "conversation-guided-001",
                "interaction_id": "interaction-guided-001",
                "parameters": {},
            },
        )
        assert started.status_code == 201
        workflow = started.json()
        assert workflow["current_step"] == "scope_type"

        answers = [
            ("scope_type", "project"),
            ("scope", "prj-0442"),
            (
                "period",
                {"start_date": "2026-08-01", "end_date": "2026-08-31"},
            ),
            ("audience", "executive_board"),
            ("purpose", "monthly_governance"),
            ("emphasis", ["schedule", "risks"]),
            ("additional_context", "Destacar o proximo marco."),
        ]
        for index, (step, value) in enumerate(answers):
            answered = client.post(
                f"/api/workflow-sessions/{workflow['session_id']}/answers",
                headers=AUTH_HEADER,
                json={
                    "step": step,
                    "value": value,
                    "revision": workflow["revision"],
                    "client_action_id": f"workflow-answer-{index:02d}",
                },
            )
            assert answered.status_code == 200, answered.text
            workflow = answered.json()

        assert workflow["status"] == "ready_for_confirmation"
        assert "Moderniza" in workflow["confirmation_summary"]

        confirmed = client.post(
            f"/api/workflow-sessions/{workflow['session_id']}/confirm",
            headers=AUTH_HEADER,
            json={
                "revision": workflow["revision"],
                "client_action_id": "workflow-confirm-001",
            },
        )
        assert confirmed.status_code == 202

        current = client.get(
            f"/api/workflow-sessions/{workflow['session_id']}",
            headers=AUTH_HEADER,
        )
        assert current.status_code == 200
        completed = current.json()
        assert completed["status"] == "completed"
        assert completed["execution"]["artifact"]["storage_provider"] == "local"

        download = client.get(
            completed["execution"]["artifact"]["download_url"],
            headers=AUTH_HEADER,
        )
        assert download.status_code == 200
        assert download.content.startswith(b"PK")


def test_cli_executes_same_local_service(monkeypatch, capsys):
    monkeypatch.setenv("REPORT_CLI_TOKEN", "report-test-token")

    result = report_cli_main(
        [
            "--scope-type",
            "project",
            "--scope-id",
            "prj-0442",
            "--start-date",
            "2026-08-01",
            "--end-date",
            "2026-08-31",
            "--emphasis",
            "risks",
        ]
    )

    output = json.loads(capsys.readouterr().out)
    assert result == 0
    assert output["status"] == "completed_with_warnings"
    assert output["artifact"]["file_name"].endswith(".docx")



def test_normalize_emphasis_maps_free_text_deterministically():
    from app.reports.schemas import Emphasis, normalize_emphasis

    assert normalize_emphasis("riscos e cronograma") == [
        Emphasis.RISKS,
        Emphasis.SCHEDULE,
    ]
    # acentos, separadores variados e deduplicacao
    assert normalize_emphasis("Benefícios; benefícios, marcos") == [
        Emphasis.BENEFITS,
        Emphasis.MILESTONES,
    ]
    # termos desconhecidos sao ignorados
    assert normalize_emphasis("qualquer coisa irrelevante") == []
    # vazio => visao geral (lista vazia)
    assert normalize_emphasis(None) == []
    assert normalize_emphasis("") == []


def test_report_options_derives_emphasis_from_free_text():
    from app.reports.schemas import Audience, Emphasis, Purpose, ReportOptions

    options = ReportOptions(
        audience=Audience.PMO,
        purpose=Purpose.MONTHLY_GOVERNANCE,
        emphasis_text="riscos, atrasos",
    )
    assert options.emphasis == [Emphasis.RISKS, Emphasis.SCHEDULE]
    # enum explicito tem precedencia sobre o texto livre
    explicit = ReportOptions(
        audience=Audience.PMO,
        purpose=Purpose.MONTHLY_GOVERNANCE,
        emphasis=[Emphasis.BENEFITS],
        emphasis_text="riscos",
    )
    assert explicit.emphasis == [Emphasis.BENEFITS]


def test_supabase_storage_keeps_object_private_and_downloads_through_backend(
    monkeypatch, tmp_path
):
    source = tmp_path / "report.docx"
    source.write_bytes(b"PK-private-docx")
    monkeypatch.setenv("SUPABASE_URL", "https://project.supabase.co")
    monkeypatch.setenv("SUPABASE_SERVICE_ROLE_KEY", "service-role-secret")
    monkeypatch.setenv("SUPABASE_STORAGE_BUCKET", "generated-reports")
    calls = []

    def fake_get(url, **kwargs):
        calls.append(("GET", url, kwargs))
        if "/storage/v1/bucket/" in url:
            return SimpleNamespace(status_code=200, content=b"", json=lambda: {})
        return SimpleNamespace(
            status_code=200, content=b"PK-private-docx", json=lambda: {}
        )

    def fake_post(url, **kwargs):
        calls.append(("POST", url, kwargs))
        return SimpleNamespace(status_code=200, content=b"", json=lambda: {})

    monkeypatch.setattr("app.reports.storage.httpx.get", fake_get)
    monkeypatch.setattr("app.reports.storage.httpx.post", fake_post)
    artifact = ReportArtifact(
        artifact_id="11111111-1111-4111-8111-111111111111",
        file_name="report.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        size_bytes=15,
        sha256="a" * 64,
        download_url="/api/report-requests/execution-1/artifact",
        storage_provider="local",
        storage_reference=str(source),
        created_at=datetime(2026, 9, 21, tzinfo=timezone.utc),
    )

    store = ArtifactStore("supabase")
    persisted = store.persist(artifact)
    downloaded = store.read(persisted, tmp_path)

    assert persisted.storage_provider == "supabase"
    assert persisted.storage_reference.startswith("2026/09/")
    assert persisted.download_url.startswith("/api/report-requests/")
    assert downloaded.data == b"PK-private-docx"
    assert not source.exists()
    assert all("service-role-secret" not in url for _, url, _ in calls)


def test_bucket_rollback_requires_server_side_credentials(monkeypatch):
    monkeypatch.delenv("SUPABASE_URL", raising=False)
    monkeypatch.delenv("SUPABASE_SERVICE_ROLE_KEY", raising=False)

    with pytest.raises(RuntimeError, match="SUPABASE_URL"):
        rollback_report_storage_bucket.delete_bucket()


def test_bucket_rollback_uses_official_storage_api(monkeypatch):
    monkeypatch.setenv("SUPABASE_URL", "https://project.supabase.co")
    monkeypatch.setenv("SUPABASE_SERVICE_ROLE_KEY", "service-role-secret")
    calls = []

    class FakeResponse:
        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

    def fake_urlopen(request, timeout):
        calls.append((request, timeout))
        return FakeResponse()

    monkeypatch.setattr(rollback_report_storage_bucket, "urlopen", fake_urlopen)

    assert rollback_report_storage_bucket.delete_bucket() is True
    request, timeout = calls[0]
    assert request.method == "DELETE"
    assert request.full_url.endswith("/storage/v1/bucket/generated-reports")
    assert request.get_header("Authorization") == "Bearer service-role-secret"
    assert timeout == 20


@pytest.mark.parametrize(
    "mutation",
    [
        lambda payload: payload.update(contract_version="9.9"),
        lambda payload: payload.update(unknown_field="nao permitido"),
    ],
)
def test_external_contract_rejects_unknown_version_and_fields(mutation):
    payload = _prepare_payload()
    mutation(payload)
    with TestClient(app) as client:
        response = _prepare(client, payload=payload)

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "REPORT_INVALID_INPUT"


@pytest.mark.parametrize(
    ("review", "summary", "expected_warning"),
    [
        (
            {
                "schema_version": "1.0",
                "approved": False,
                "numeric_mismatches": ["percentual divergente"],
            },
            {"schema_version": "1.0", "summary": "Resumo nao confiavel"},
            "nao passou na validacao",
        ),
        (
            {"schema_version": "1.0", "approved": True},
            {
                "schema_version": "1.0",
                "summary": "Resumo com fonte inventada",
                "citation_source_ids": ["fonte-inexistente"],
            },
            "citacoes desconhecidas",
        ),
    ],
)
def test_untrusted_ai_builder_output_falls_back_to_deterministic_text(
    monkeypatch, review, summary, expected_warning
):
    monkeypatch.setenv("REPORT_EXECUTION_MODE", "power_automate")
    monkeypatch.setattr(
        "app.reports.service.ReportService.dispatch_power_automate",
        lambda _service, _execution_id: None,
    )
    _reset_runtime()
    with TestClient(app) as client:
        prepared = _prepare(client).json()
        created = client.post(
            "/api/report-requests", json=_create_payload(prepared), headers=AUTH_HEADER
        ).json()
        callback = client.post(
            "/api/integrations/power-automate/report-results",
            headers=CALLBACK_HEADER,
            json={
                "contract_version": "1.0",
                "callback_id": f"callback-{expected_warning[:8]}",
                "execution_id": created["execution_id"],
                "correlation_id": created["correlation_id"],
                "flow_run_id": "flow-ai-review",
                "status": "completed",
                "result": {
                    "schema_version": "1.0",
                    "foundation_review": review,
                    "executive_summary": summary,
                },
            },
        )
        downloaded = client.get(
            f"/api/report-requests/{created['execution_id']}/artifact",
            headers=AUTH_HEADER,
        )

    assert callback.status_code == 200
    assert any(expected_warning in item for item in callback.json()["warnings"])
    document = Document(BytesIO(downloaded.content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert summary["summary"] not in text


@pytest.mark.parametrize("status", ["failed", "blocked"])
def test_failure_callbacks_preserve_controlled_error_without_artifact(monkeypatch, status):
    monkeypatch.setenv("REPORT_EXECUTION_MODE", "power_automate")
    monkeypatch.setattr(
        "app.reports.service.ReportService.dispatch_power_automate",
        lambda _service, _execution_id: None,
    )
    _reset_runtime()
    with TestClient(app) as client:
        prepared = _prepare(client).json()
        created = client.post(
            "/api/report-requests", json=_create_payload(prepared), headers=AUTH_HEADER
        ).json()
        callback = client.post(
            "/api/integrations/power-automate/report-results",
            headers=CALLBACK_HEADER,
            json={
                "contract_version": "1.0",
                "callback_id": f"callback-{status}-001",
                "execution_id": created["execution_id"],
                "correlation_id": created["correlation_id"],
                "flow_run_id": f"flow-{status}",
                "status": status,
                "error_code": "FLOW_CONTROLLED_FAILURE",
                "error_message": "Falha controlada do fluxo.",
            },
        )

    assert callback.status_code == 200
    assert callback.json()["status"] == status
    assert callback.json()["artifact"] is None
    assert callback.json()["error"] == {
        "code": "FLOW_CONTROLLED_FAILURE",
        "message": "Falha controlada do fluxo.",
    }


def test_dispatch_retries_with_idempotency_and_accepts_late_callback(monkeypatch):
    monkeypatch.setenv("REPORT_EXECUTION_MODE", "power_automate")
    monkeypatch.setenv("POWER_AUTOMATE_REPORT_TRIGGER_URL", "https://flow.test/run")
    monkeypatch.setenv("POWER_AUTOMATE_TRIGGER_TOKEN", "flow-token")
    monkeypatch.setenv("REPORT_CALLBACK_BASE_URL", "https://api.test")
    monkeypatch.setenv("POWER_AUTOMATE_DISPATCH_ATTEMPTS", "3")
    calls = []

    def timeout(url, **kwargs):
        calls.append((url, kwargs))
        raise httpx.ReadTimeout("sem confirmacao", request=httpx.Request("POST", url))

    monkeypatch.setattr("app.reports.service.httpx.post", timeout)
    _reset_runtime()
    with TestClient(app) as client:
        prepared = _prepare(client).json()
        created = client.post(
            "/api/report-requests", json=_create_payload(prepared), headers=AUTH_HEADER
        ).json()
        blocked = client.get(
            f"/api/report-requests/{created['execution_id']}", headers=AUTH_HEADER
        )
        recovered = client.post(
            "/api/integrations/power-automate/report-results",
            headers=CALLBACK_HEADER,
            json={
                "contract_version": "1.0",
                "callback_id": "callback-late-001",
                "execution_id": created["execution_id"],
                "correlation_id": created["correlation_id"],
                "flow_run_id": "flow-late-001",
                "status": "completed",
                "result": {"schema_version": "1.0"},
            },
        )

    assert len(calls) == 3
    assert all(call[1]["headers"]["Idempotency-Key"] for call in calls)
    assert blocked.json()["status"] == "blocked"
    assert blocked.json()["error"]["code"] == "REPORT_AUTOMATION_DISPATCH_UNCERTAIN"
    assert recovered.status_code == 200
    assert recovered.json()["status"] == "completed"
