"""Renderizacao DOCX do relatorio executivo no modo local."""

from __future__ import annotations

import hashlib
import os
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4

from app.reports.errors import ReportConfigurationError
from app.reports.schemas import PreparedReport, ReportArtifact, ReportAutomationResult

DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


class DocxReportRenderer:
    def __init__(self, artifact_directory: Path | None = None) -> None:
        configured = os.getenv("REPORT_ARTIFACTS_DIR", "").strip()
        self.artifact_directory = artifact_directory or (
            Path(configured) if configured else Path.cwd() / ".runtime" / "reports"
        )

    def render(
        self,
        execution_id: str,
        prepared: PreparedReport,
        snapshot: dict[str, Any],
        automation_result: ReportAutomationResult | None = None,
    ) -> ReportArtifact:
        try:
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.shared import Pt
        except ImportError as error:
            raise ReportConfigurationError() from error

        output_directory = self.artifact_directory.resolve()
        output_directory.mkdir(parents=True, exist_ok=True)
        safe_title = _safe_file_part(prepared.scope_title)
        file_name = f"relatorio-executivo-{safe_title}-{execution_id}.docx"
        output_path = (output_directory / file_name).resolve()
        if output_directory not in output_path.parents:
            raise ReportConfigurationError()

        document = Document()
        styles = document.styles
        styles["Normal"].font.name = "Arial"
        styles["Normal"].font.size = Pt(10)

        title = document.add_heading("Relatório executivo", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle = document.add_paragraph(prepared.scope_title)
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.add_paragraph(
            f"Período: {_period_label(prepared)} | "
            f"Referência dos dados: {_format_datetime(prepared.data_reference_at)}"
        )

        _add_section(document, "1. Síntese executiva")
        document.add_paragraph(
            automation_result.executive_summary.summary
            if automation_result and automation_result.executive_summary
            else _executive_summary(snapshot)
        )

        _add_section(document, "2. Situação e indicadores")
        _add_project_status_table(document, snapshot["projects"])

        _add_section(document, "3. Previsto versus realizado")
        _add_progress_table(document, snapshot["projects"])

        _add_section(document, "4. Marcos e próximas entregas")
        _add_items_table(
            document,
            snapshot["milestones"],
            ["Marco", "Data prevista", "Situação"],
            ["description", "planned_date", "status"],
        )

        _add_section(document, "5. Riscos e pontos de atenção")
        _add_items_table(
            document,
            snapshot["risks"],
            ["Risco ou problema", "Criticidade", "Tratamento", "Responsável"],
            ["title", "criticality", "mitigation", "owner"],
        )

        _add_section(document, "6. Realizações, mudanças e decisões")
        _add_achievements(document, snapshot["projects"])
        _add_items_table(
            document,
            snapshot["changes"],
            ["Mudança", "Situação", "Impacto"],
            ["title", "status", "impact"],
        )

        _add_section(document, "7. Benefícios e alinhamento estratégico")
        document.add_paragraph(snapshot["subportfolio"]["strategic_objective"])
        _add_items_table(
            document,
            snapshot["benefits"],
            ["Benefício", "Indicador", "Valor esperado", "Início da captura"],
            ["description", "indicator", "expected_value", "capture_start"],
        )

        additional_context = prepared.options.additional_context
        if automation_result and automation_result.normalized_context:
            additional_context = automation_result.normalized_context.normalized_text
        if additional_context:
            _add_section(document, "8. Contexto adicional confirmado pelo usuário")
            document.add_paragraph(additional_context)

        _add_section(document, "Fontes, atualização e limitações")
        document.add_paragraph(
            "Este documento foi gerado automaticamente a partir de uma fotografia "
            "imutável dos dados autorizados para o período confirmado."
        )
        for source in snapshot["sources"]:
            document.add_paragraph(
                f"{source['title']} — {source['locator']} — atualização: "
                f"{_format_datetime(source['last_updated'])}",
                style="List Bullet",
            )
        for limitation in snapshot.get("limitations", []):
            document.add_paragraph(f"Limitação: {limitation}", style="List Bullet")

        document.add_paragraph(
            f"Identificadores: execução {execution_id}; correlação "
            f"{prepared.correlation_id}; contrato {prepared.contract_version}."
        )
        document.save(output_path)

        content = output_path.read_bytes()
        return ReportArtifact(
            artifact_id=str(uuid4()),
            file_name=file_name,
            content_type=DOCX_CONTENT_TYPE,
            size_bytes=len(content),
            sha256=hashlib.sha256(content).hexdigest(),
            download_url=f"/api/report-requests/{execution_id}/artifact",
            storage_provider="local",
            storage_reference=str(output_path),
            created_at=datetime.now(timezone.utc),
        )


def _add_section(document: Any, title: str) -> None:
    document.add_heading(title, level=1)


def _add_project_status_table(document: Any, projects: list[dict[str, Any]]) -> None:
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for cell, label in zip(
        table.rows[0].cells,
        ["Projeto", "Situação", "Fase", "Responsável", "Atualização"],
        strict=True,
    ):
        cell.text = label
    for project in projects:
        cells = table.add_row().cells
        values = [
            f"{project['code']} — {project['title']}",
            project["status"],
            project["phase"],
            project["manager"],
            project["updated_at"],
        ]
        for cell, value in zip(cells, values, strict=True):
            cell.text = str(value)


def _add_progress_table(document: Any, projects: list[dict[str, Any]]) -> None:
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, label in zip(
        table.rows[0].cells,
        ["Projeto", "Previsto (%)", "Realizado (%)", "Desvio (p.p.)"],
        strict=True,
    ):
        cell.text = label
    for project in projects:
        planned = float(project["planned_progress"])
        actual = float(project["actual_progress"])
        values = [
            project["code"],
            f"{planned:.1f}",
            f"{actual:.1f}",
            f"{actual - planned:.1f}",
        ]
        cells = table.add_row().cells
        for cell, value in zip(cells, values, strict=True):
            cell.text = str(value)


def _add_items_table(
    document: Any,
    items: list[dict[str, Any]],
    headers: list[str],
    fields: list[str],
) -> None:
    if not items:
        document.add_paragraph("Não há registros disponíveis para este recorte.")
        return
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, headers, strict=True):
        cell.text = label
    for item in items:
        cells = table.add_row().cells
        for cell, field in zip(cells, fields, strict=True):
            cell.text = str(item.get(field) or "Não informado")


def _add_achievements(document: Any, projects: list[dict[str, Any]]) -> None:
    for project in projects:
        document.add_paragraph(
            f"{project['code']}: {project.get('achievements') or 'Não informado'}",
            style="List Bullet",
        )


def _executive_summary(snapshot: dict[str, Any]) -> str:
    projects = snapshot["projects"]
    actual = sum(float(item["actual_progress"]) for item in projects) / len(projects)
    planned = sum(float(item["planned_progress"]) for item in projects) / len(projects)
    critical = sum(
        1
        for item in snapshot["risks"]
        if str(item.get("criticality", "")).lower() in {"alta", "critica", "crítica"}
    )
    return (
        f"O recorte reúne {len(projects)} projeto(s), com avanço físico médio "
        f"realizado de {actual:.1f}% frente a {planned:.1f}% previsto "
        f"({actual - planned:+.1f} ponto(s) percentual(is)). "
        f"Foram identificados {critical} risco(s) ou problema(s) de alta "
        "criticidade. As interpretações deste resumo derivam somente dos "
        "valores presentes na fotografia de dados indicada ao final do documento."
    )


def _period_label(prepared: PreparedReport) -> str:
    if prepared.period.label:
        return prepared.period.label
    return (
        f"{prepared.period.start_date:%d/%m/%Y} a {prepared.period.end_date:%d/%m/%Y}"
    )


def _format_datetime(value: datetime | str) -> str:
    if isinstance(value, str):
        value = datetime.fromisoformat(value.replace("Z", "+00:00"))
    return value.astimezone().strftime("%d/%m/%Y %H:%M %Z")


def _safe_file_part(value: str) -> str:
    normalized = re.sub(r"[^a-zA-Z0-9_-]+", "-", value).strip("-").lower()
    return normalized[:80] or "escopo"
